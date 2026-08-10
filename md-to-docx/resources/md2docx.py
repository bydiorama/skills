#!/usr/bin/env python3
"""md2docx.py — convert Markdown to DOCX using python-docx.

Supports:
  - H1–H6 headings (# .. ######), navy, sized 22/16/13/12/11/10 pt
  - GFM tables with pipe-aware tokenization (backticks and \\| escapes respected)
  - Bullet lists (-, *, +) and numbered lists (1.) with 2-space-per-level nesting
  - Inline: ***bold-italic***, **bold**, *italic*, `code`, [text](url) hyperlinks
  - Blockquotes (>), horizontal rules (--- / *** / ___), task checkboxes (- [x] / - [ ])
  - Fenced code blocks (``` and ~~~) rendered as Consolas 10 pt with line breaks
  - YAML frontmatter (skipped between leading --- fences)
  - UTF-8 BOM (read via utf-8-sig)

No pandoc, no network, idempotent (overwrites output).
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

try:
    from docx import Document
    from docx.enum.table import WD_ALIGN_VERTICAL
    from docx.opc.constants import RELATIONSHIP_TYPE
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    from docx.shared import Pt, RGBColor, Inches
except ImportError:
    sys.stderr.write(
        "python-docx is required. Install with: python3 -m pip install --user python-docx\n"
    )
    sys.exit(1)


NAVY = RGBColor(0x0B, 0x1F, 0x4F)
DARK_NAVY_HEX = "081633"
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GREY = RGBColor(0x55, 0x55, 0x55)

HEADING_SIZES = {1: 22, 2: 16, 3: 13, 4: 12, 5: 11, 6: 10}

INLINE_PATTERN = re.compile(
    r"\*\*\*(?P<bolditalic>[^*]+)\*\*\*"
    r"|\*\*(?P<bold>[^*]+)\*\*"
    r"|\*(?P<italic>[^*]+)\*"
    r"|`(?P<code>[^`]+)`"
    r"|\[(?P<link_text>[^\]]+)\]\((?P<link_url>[^)\s]+)\)"
)

FENCE_RE = re.compile(r"^(\s*)(```|~~~)\s*([\w+-]*)\s*$")
HEADING_RE = re.compile(r"^(#{1,6})\s+(.*)$")
BULLET_RE = re.compile(r"^(\s*)([-*+])\s+(.*)$")
NUMBERED_RE = re.compile(r"^(\s*)\d+\.\s+(.*)$")
HR_RE = re.compile(r"^\s*(-{3,}|_{3,}|\*{3,})\s*$")
BLOCKQUOTE_RE = re.compile(r"^\s*>\s?(.*)$")
TABLE_LINE_RE = re.compile(r"^\s*\|")
SEP_CELL_RE = re.compile(r":?-{3,}:?")


def shade_cell(cell, hex_color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.append(shd)


def add_hyperlink(paragraph, text: str, url: str) -> None:
    """Attach a Word hyperlink to an existing paragraph."""
    part = paragraph.part
    try:
        r_id = part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    except Exception:
        paragraph.add_run(f"[{text}]({url})")
        return
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    new_run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0B57D0")
    rPr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rPr.append(underline)
    new_run.append(rPr)
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    new_run.append(t)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def add_runs_with_inline(paragraph, text: str) -> None:
    """Parse inline markdown formatting and append runs to paragraph."""
    idx = 0
    for m in INLINE_PATTERN.finditer(text):
        if m.start() > idx:
            paragraph.add_run(text[idx:m.start()])
        if m.group("bolditalic") is not None:
            r = paragraph.add_run(m.group("bolditalic"))
            r.bold = True
            r.italic = True
        elif m.group("bold") is not None:
            r = paragraph.add_run(m.group("bold"))
            r.bold = True
        elif m.group("italic") is not None:
            r = paragraph.add_run(m.group("italic"))
            r.italic = True
        elif m.group("code") is not None:
            r = paragraph.add_run(m.group("code"))
            r.font.name = "Consolas"
            r.font.size = Pt(10)
        elif m.group("link_text") is not None:
            add_hyperlink(paragraph, m.group("link_text"), m.group("link_url"))
        idx = m.end()
    if idx < len(text):
        paragraph.add_run(text[idx:])


def style_heading(paragraph, level: int) -> None:
    size = HEADING_SIZES.get(level, 11)
    for run in paragraph.runs:
        run.font.size = Pt(size)
        run.font.color.rgb = NAVY
        run.bold = True


def add_horizontal_rule(doc) -> None:
    p = doc.add_paragraph()
    p_pr = p._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "auto")
    pbdr.append(bottom)
    p_pr.append(pbdr)


def tokenize_table_row(line: str) -> list[str]:
    """Split a table row on unescaped, non-backticked pipes. Strip outer pipes."""
    s = line.strip()
    if s.startswith("|"):
        s = s[1:]
    if s.endswith("|") and not s.endswith("\\|"):
        s = s[:-1]
    cells: list[str] = []
    buf: list[str] = []
    in_code = False
    i = 0
    while i < len(s):
        c = s[i]
        if c == "\\" and i + 1 < len(s) and s[i + 1] == "|":
            buf.append("|")
            i += 2
            continue
        if c == "`":
            in_code = not in_code
            buf.append(c)
            i += 1
            continue
        if c == "|" and not in_code:
            cells.append("".join(buf).strip())
            buf = []
            i += 1
            continue
        buf.append(c)
        i += 1
    cells.append("".join(buf).strip())
    return cells


def is_separator_row(cells: list[str]) -> bool:
    if not cells:
        return False
    return all(SEP_CELL_RE.fullmatch(c) is not None for c in cells)


def parse_table_rows(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    rows: list[list[str]] = []
    i = start
    while i < len(lines):
        line = lines[i]
        if not TABLE_LINE_RE.match(line):
            break
        rows.append(tokenize_table_row(line))
        i += 1
    return rows, i


def add_table(doc, rows: list[list[str]]) -> None:
    if not rows:
        return
    header = rows[0]
    body = rows[1:]
    cols = max(len(header), max((len(r) for r in body), default=0))
    if cols == 0:
        return
    table = doc.add_table(rows=1 + len(body), cols=cols)
    try:
        table.style = "Light Grid Accent 1"
    except KeyError:
        pass
    hdr_cells = table.rows[0].cells
    for j in range(cols):
        cell = hdr_cells[j]
        cell.text = ""
        p = cell.paragraphs[0]
        add_runs_with_inline(p, header[j] if j < len(header) else "")
        for run in p.runs:
            run.bold = True
            run.font.color.rgb = WHITE
        shade_cell(cell, DARK_NAVY_HEX)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    for i, row in enumerate(body, start=1):
        cells = table.rows[i].cells
        for j in range(cols):
            cell = cells[j]
            cell.text = ""
            p = cell.paragraphs[0]
            add_runs_with_inline(p, row[j] if j < len(row) else "")
    doc.add_paragraph()


def render_task_prefix(text: str) -> tuple[str, str]:
    m = re.match(r"\[([ xX])\]\s+(.*)", text)
    if not m:
        return "", text
    return ("☒ " if m.group(1).lower() == "x" else "☐ "), m.group(2)


def add_fenced_code(doc, code_lines: list[str]) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    for idx, ln in enumerate(code_lines):
        r = p.add_run(ln)
        r.font.name = "Consolas"
        r.font.size = Pt(10)
        if idx < len(code_lines) - 1:
            br = OxmlElement("w:br")
            r._r.append(br)


def skip_frontmatter(lines: list[str]) -> int:
    """If the file starts with YAML frontmatter (--- .. ---), return the index
    after the closing fence. Otherwise return 0.
    """
    i = 0
    while i < len(lines) and not lines[i].strip():
        i += 1
    if i >= len(lines) or lines[i].strip() != "---":
        return 0
    j = i + 1
    while j < len(lines):
        if lines[j].strip() == "---":
            return j + 1
        j += 1
    return 0  # no closing fence → treat as normal content


def numbered_list_style(level: int) -> str:
    return {1: "List Number", 2: "List Number 2", 3: "List Number 3"}.get(level, "List Number")


def bullet_list_style(level: int) -> str:
    return {1: "List Bullet", 2: "List Bullet 2", 3: "List Bullet 3"}.get(level, "List Bullet")


def indent_to_level(indent: int) -> int:
    return 1 + (indent // 2)


def convert(input_path: Path, output_path: Path) -> None:
    # utf-8-sig strips a leading BOM if present, passes through otherwise.
    text = input_path.read_text(encoding="utf-8-sig")
    lines = text.splitlines()

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    i = skip_frontmatter(lines)

    in_fence = False
    fence_marker: str | None = None
    fence_buf: list[str] = []

    while i < len(lines):
        raw = lines[i]
        # Normalize tabs in leading whitespace so indent arithmetic is stable.
        leading_tabs = re.match(r"^\t+", raw)
        if leading_tabs:
            raw_indent_normalized = "  " * len(leading_tabs.group(0)) + raw[len(leading_tabs.group(0)):]
        else:
            raw_indent_normalized = raw
        line = raw_indent_normalized.rstrip()

        # Fenced-code-block state machine.
        fence_m = FENCE_RE.match(raw)
        if in_fence:
            if fence_m and fence_m.group(2) == fence_marker:
                add_fenced_code(doc, fence_buf)
                in_fence = False
                fence_marker = None
                fence_buf = []
                i += 1
                continue
            fence_buf.append(raw)
            i += 1
            continue
        if fence_m:
            in_fence = True
            fence_marker = fence_m.group(2)
            fence_buf = []
            i += 1
            continue

        if not line.strip():
            i += 1
            continue

        if HR_RE.match(line):
            add_horizontal_rule(doc)
            i += 1
            continue

        m = HEADING_RE.match(line)
        if m:
            level = len(m.group(1))
            heading_text = m.group(2).strip()
            try:
                p = doc.add_heading(level=min(level, 4))
            except Exception:
                p = doc.add_paragraph()
            add_runs_with_inline(p, heading_text)
            style_heading(p, level)
            i += 1
            continue

        bq = BLOCKQUOTE_RE.match(line)
        if bq:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            add_runs_with_inline(p, bq.group(1))
            for run in p.runs:
                run.italic = True
                run.font.color.rgb = GREY
            i += 1
            continue

        if TABLE_LINE_RE.match(line):
            rows, next_i = parse_table_rows(lines, i)
            if len(rows) >= 2 and is_separator_row(rows[1]):
                header = rows[0]
                body = [r for r in rows[2:] if not is_separator_row(r)]
                add_table(doc, [header] + body)
                i = next_i
                continue

        bullet_m = BULLET_RE.match(line)
        if bullet_m:
            indent = len(bullet_m.group(1))
            content = bullet_m.group(3)
            prefix, content = render_task_prefix(content)
            level = min(indent_to_level(indent), 3)
            try:
                p = doc.add_paragraph(style=bullet_list_style(level))
            except KeyError:
                p = doc.add_paragraph()
            if prefix:
                p.add_run(prefix)
            add_runs_with_inline(p, content)
            i += 1
            continue

        num_m = NUMBERED_RE.match(line)
        if num_m:
            indent = len(num_m.group(1))
            content = num_m.group(2)
            level = min(indent_to_level(indent), 3)
            try:
                p = doc.add_paragraph(style=numbered_list_style(level))
            except KeyError:
                p = doc.add_paragraph()
            add_runs_with_inline(p, content)
            i += 1
            continue

        # Default paragraph — merge consecutive non-blank, non-structural lines.
        para_lines = [line]
        j = i + 1
        while j < len(lines):
            nxt_raw = lines[j]
            nxt = nxt_raw.rstrip()
            if not nxt.strip():
                break
            if (
                FENCE_RE.match(nxt_raw)
                or HEADING_RE.match(nxt)
                or HR_RE.match(nxt)
                or BLOCKQUOTE_RE.match(nxt)
                or TABLE_LINE_RE.match(nxt)
                or BULLET_RE.match(nxt)
                or NUMBERED_RE.match(nxt)
            ):
                break
            para_lines.append(nxt)
            j += 1
        p = doc.add_paragraph()
        add_runs_with_inline(p, " ".join(para_lines))
        i = j

    # Unterminated fence: flush what we gathered so content isn't lost.
    if in_fence and fence_buf:
        add_fenced_code(doc, fence_buf)

    doc.save(str(output_path))


def main() -> int:
    if len(sys.argv) < 2:
        print("Usage: md2docx.py <input.md> [output.docx]", file=sys.stderr)
        return 2
    input_path = Path(sys.argv[1]).expanduser()
    if not input_path.exists():
        print(f"Input not found: {input_path}", file=sys.stderr)
        return 2
    if input_path.suffix.lower() not in {".md", ".markdown"}:
        print(f"Input must be .md or .markdown: {input_path}", file=sys.stderr)
        return 2
    if len(sys.argv) >= 3:
        output_path = Path(sys.argv[2]).expanduser()
    else:
        output_path = input_path.with_suffix(".docx")
    convert(input_path, output_path)
    size = output_path.stat().st_size
    print(f"Wrote {output_path} ({size:,} bytes)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
